#!/usr/bin/env python3
"""Inspect a PPTX template and generate manifest, roles, and cover preview.

Usage:
    python native_template_inspect.py <template.pptx> -o <template_dir> --name <display_name>
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

# Ensure the scripts directory is on sys.path so sibling modules can be found
_scripts_dir = str(Path(__file__).resolve().parent)
if _scripts_dir not in sys.path:
    sys.path.insert(0, _scripts_dir)


def _infer_roles(page_count: int) -> dict[str, int]:
    """Infer default role-to-slide mapping based on page count."""
    if page_count <= 0:
        return {"cover": 1, "toc": 1, "content": 1, "thanks": 1}
    if page_count == 1:
        return {"cover": 1, "toc": 1, "content": 1, "thanks": 1}
    if page_count == 2:
        return {"cover": 1, "toc": 2, "content": 2, "thanks": 2}
    if page_count == 3:
        return {"cover": 1, "toc": 2, "content": 2, "thanks": 3}
    return {"cover": 1, "toc": 2, "content": 3, "thanks": 4}


def _default_zones() -> dict[str, dict[str, dict[str, float]]]:
    """Return default zone layout."""
    return {
        "cover": {
            "title": {"x": 0.08, "y": 0.20, "w": 0.84, "h": 0.18},
            "subtitle": {"x": 0.10, "y": 0.42, "w": 0.80, "h": 0.10},
        },
        "toc": {
            "title": {"x": 0.06, "y": 0.06, "w": 0.88, "h": 0.10},
            "body": {"x": 0.12, "y": 0.22, "w": 0.76, "h": 0.62},
        },
        "content": {
            "title": {"x": 0.06, "y": 0.06, "w": 0.88, "h": 0.10},
            "body": {"x": 0.08, "y": 0.20, "w": 0.84, "h": 0.68},
        },
        "thanks": {
            "title": {"x": 0.10, "y": 0.34, "w": 0.80, "h": 0.16},
            "subtitle": {"x": 0.16, "y": 0.52, "w": 0.68, "h": 0.10},
        },
    }


def _detect_zones_for_slide(shapes: list[dict]) -> dict[str, dict[str, float]]:
    """Auto-detect title and body zones from a slide's shapes."""
    text_shapes = [s for s in shapes if s.get("has_text")]

    if not text_shapes:
        return {
            "title": {"x": 0.06, "y": 0.06, "w": 0.88, "h": 0.10},
            "body": {"x": 0.08, "y": 0.20, "w": 0.84, "h": 0.68},
        }

    # Find topmost text shape as title
    title_shape = min(text_shapes, key=lambda s: s["y"])

    # Find largest remaining text shape as body
    remaining = [s for s in text_shapes if s is not title_shape]
    if remaining:
        body_shape = max(remaining, key=lambda s: s["w"] * s["h"])
    else:
        body_shape = None

    zones: dict[str, dict[str, float]] = {
        "title": {
            "x": round(title_shape["x"], 3),
            "y": round(title_shape["y"], 3),
            "w": round(title_shape["w"], 3),
            "h": round(title_shape["h"], 3),
        },
    }

    if body_shape:
        zones["body"] = {
            "x": round(body_shape["x"], 3),
            "y": round(body_shape["y"], 3),
            "w": round(body_shape["w"], 3),
            "h": round(body_shape["h"], 3),
        }

    return zones


def inspect_pptx(pptx_path: Path, output_dir: Path, display_name: str) -> int:
    """Inspect PPTX and generate manifest, roles, and cover preview."""
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    slide_count = len(prs.slides)
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    # Determine canvas format
    ratio = slide_width / slide_height if slide_height else 1.777
    canvas_format = "ppt43" if abs(ratio - 4 / 3) < 0.1 else "ppt169"

    # Inspect each slide
    slides_data: list[dict] = []
    for i, slide in enumerate(prs.slides):
        shapes_data: list[dict] = []
        for shape in slide.shapes:
            shape_info = {
                "shape_id": shape.shape_id,
                "name": shape.name,
                "has_text": shape.has_text_frame,
                "x": round(shape.left / slide_width, 3) if slide_width else 0,
                "y": round(shape.top / slide_height, 3) if slide_height else 0,
                "w": round(shape.width / slide_width, 3) if slide_width else 0,
                "h": round(shape.height / slide_height, 3) if slide_height else 0,
            }
            if shape.has_text_frame:
                shape_info["text"] = shape.text_frame.text[:200]
            shapes_data.append(shape_info)
        slides_data.append({"slide_index": i, "shapes": shapes_data})

    # Generate manifest
    manifest = {
        "slide_count": slide_count,
        "slide_width_emu": slide_width,
        "slide_height_emu": slide_height,
        "canvas_format": canvas_format,
        "slides": slides_data,
    }

    # Generate roles with auto-detected zones
    roles = _infer_roles(slide_count)
    zones = _default_zones()

    # Override zones with auto-detected ones from the actual slides
    for role_name, slide_idx in roles.items():
        slide_idx_0 = slide_idx - 1  # Convert to 0-based
        if 0 <= slide_idx_0 < len(slides_data):
            detected = _detect_zones_for_slide(slides_data[slide_idx_0]["shapes"])
            if detected:
                zones[role_name] = detected

    roles_data = {
        "version": 1,
        "roles": roles,
        "zones": zones,
    }

    # Create output directory
    output_dir.mkdir(parents=True, exist_ok=True)

    # Copy template PPTX
    shutil.copy2(pptx_path, output_dir / "template.pptx")

    # Write manifest
    (output_dir / "template_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    # Write roles
    (output_dir / "template_roles.json").write_text(
        json.dumps(roles_data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    # Generate cover preview
    cover_path = output_dir / "01_cover"
    cover_png = output_dir / "01_cover.png"
    try:
        from pptx_render_slide import render_first_slide

        render_first_slide(pptx_path, cover_path)
    except Exception:
        pass

    # If render_first_slide didn't produce a PNG, try PIL fallback
    if not cover_png.exists():
        try:
            from PIL import Image, ImageDraw, ImageFont

            img = Image.new("RGB", (1280, 720), "white")
            draw = ImageDraw.Draw(img)
            draw.rectangle([(2, 2), (1277, 717)], outline="#CCCCCC", width=2)
            try:
                font = ImageFont.load_default(size=36)
            except Exception:
                font = ImageFont.load_default()
            bbox = draw.textbbox((0, 0), display_name, font=font)
            text_w = bbox[2] - bbox[0]
            text_h = bbox[3] - bbox[1]
            draw.text(
                ((1280 - text_w) // 2, (720 - text_h) // 2),
                display_name,
                fill="#666666",
                font=font,
            )
            img.save(str(cover_png))
        except Exception:
            pass  # Non-critical: cover preview is optional

    print(f"[OK] Native template inspected: {output_dir}")
    print(f"     Slides: {slide_count}")
    print(f"     Canvas: {canvas_format}")
    print(f"     Roles: {roles}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect PPTX template for native mode")
    parser.add_argument("pptx_path", help="Path to the PPTX template file")
    parser.add_argument("-o", "--output", required=True, help="Output template directory")
    parser.add_argument("--name", required=True, help="Display name for the template")
    args = parser.parse_args()

    pptx_path = Path(args.pptx_path)
    output_dir = Path(args.output)

    if not pptx_path.exists():
        print(f"Error: PPTX file not found: {pptx_path}", file=sys.stderr)
        return 1

    return inspect_pptx(pptx_path, output_dir, args.name)


if __name__ == "__main__":
    raise SystemExit(main())
