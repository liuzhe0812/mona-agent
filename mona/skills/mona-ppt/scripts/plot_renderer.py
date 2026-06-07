#!/usr/bin/env python3
"""Render Matplotlib-style plot assets for mona-ppt decks.

This utility generates chart assets, not full slide SVG pages. The main SVG
page authoring flow can embed the produced SVG/PNG into a richer layout.
"""

from __future__ import annotations

import argparse
import csv
import html
import json
from pathlib import Path
from typing import Any


MATLAB_COLORS = [
    "#0072BD",
    "#D95319",
    "#EDB120",
    "#7E2F8E",
    "#77AC30",
    "#4DBEEE",
    "#A2142F",
]


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _rows_from_json(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, dict):
        if isinstance(payload.get("data"), list):
            return _rows_from_json(payload["data"])
        if payload and all(isinstance(v, list) for v in payload.values()):
            keys = list(payload.keys())
            length = max(len(payload[k]) for k in keys)
            rows: list[dict[str, Any]] = []
            for i in range(length):
                rows.append({k: payload[k][i] if i < len(payload[k]) else None for k in keys})
            return rows
        return [payload]
    if isinstance(payload, list):
        rows = []
        for item in payload:
            if isinstance(item, dict):
                rows.append(item)
            elif isinstance(item, list):
                rows.append({str(i + 1): value for i, value in enumerate(item)})
        return rows
    raise ValueError("JSON must be an object, an array of objects, or a column-list object")


def load_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    if suffix == ".json":
        return _rows_from_json(json.loads(path.read_text(encoding="utf-8")))
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX input requires openpyxl") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(v) if v is not None else f"col{i + 1}" for i, v in enumerate(values[0])]
        return [dict(zip(headers, row)) for row in values[1:]]
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("Legacy XLS input requires xlrd; resave as .xlsx if xlrd is unavailable") from exc
        workbook = xlrd.open_workbook(str(path))
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows == 0:
            return []
        headers = [
            str(sheet.cell_value(0, col)) if sheet.cell_value(0, col) not in (None, "") else f"col{col + 1}"
            for col in range(sheet.ncols)
        ]
        rows = []
        for row_idx in range(1, sheet.nrows):
            rows.append({
                headers[col]: sheet.cell_value(row_idx, col)
                for col in range(sheet.ncols)
            })
        return rows
    raise ValueError(f"Unsupported input type: {suffix}")


def _headers(rows: list[dict[str, Any]]) -> list[str]:
    seen: list[str] = []
    for row in rows:
        for key in row:
            if key not in seen:
                seen.append(key)
    return seen


def _numeric_values(rows: list[dict[str, Any]], column: str) -> list[float | None]:
    return [_to_float(row.get(column)) for row in rows]


def _choose_y_columns(rows: list[dict[str, Any]], requested: list[str]) -> list[str]:
    headers = _headers(rows)
    if requested:
        missing = [col for col in requested if col not in headers]
        if missing:
            raise ValueError(f"Unknown y column(s): {', '.join(missing)}")
        return requested
    numeric = [
        col for col in headers
        if sum(value is not None for value in _numeric_values(rows, col)) >= 2
    ]
    if not numeric:
        raise ValueError("No numeric columns found for plotting")
    return numeric[: min(3, len(numeric))]


def _choose_x_values(
    rows: list[dict[str, Any]],
    requested_x: str | None,
    y_columns: list[str],
) -> tuple[str, list[float]]:
    headers = _headers(rows)
    if requested_x:
        if requested_x not in headers:
            raise ValueError(f"Unknown x column: {requested_x}")
        values = _numeric_values(rows, requested_x)
        if any(value is None for value in values):
            raise ValueError(f"x column must be numeric: {requested_x}")
        return requested_x, [float(value) for value in values if value is not None]

    for col in headers:
        if col in y_columns:
            continue
        values = _numeric_values(rows, col)
        if len(values) == len(rows) and all(value is not None for value in values):
            return col, [float(value) for value in values if value is not None]
    return "index", list(range(1, len(rows) + 1))


def _nice_range(values: list[float]) -> tuple[float, float]:
    min_value = min(values)
    max_value = max(values)
    if min_value == max_value:
        pad = abs(min_value) * 0.1 or 1
        return min_value - pad, max_value + pad
    pad = (max_value - min_value) * 0.08
    return min_value - pad, max_value + pad


def _format_tick(value: float) -> str:
    if abs(value) >= 1000 or (0 < abs(value) < 0.01):
        return f"{value:.2e}"
    if float(value).is_integer():
        return str(int(value))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def _render_svg_fallback(
    output_path: Path,
    *,
    kind: str,
    x_name: str,
    x_values: list[float],
    series: list[tuple[str, list[float]]],
    title: str | None,
    x_label: str | None,
    y_label: str | None,
    width: float,
    height: float,
    dpi: int,
) -> None:
    svg_width = int(width * dpi)
    svg_height = int(height * dpi)
    margin_left = 78
    margin_right = 42
    margin_top = 62 if title else 42
    margin_bottom = 64
    plot_width = svg_width - margin_left - margin_right
    plot_height = svg_height - margin_top - margin_bottom

    y_all = [value for _, values in series for value in values]
    x_min, x_max = _nice_range(x_values)
    y_min, y_max = _nice_range(y_all)

    def sx(value: float) -> float:
        return margin_left + (value - x_min) / (x_max - x_min) * plot_width

    def sy(value: float) -> float:
        return margin_top + plot_height - (value - y_min) / (y_max - y_min) * plot_height

    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{svg_width}" height="{svg_height}" viewBox="0 0 {svg_width} {svg_height}">',
        '<rect width="100%" height="100%" fill="#FFFFFF"/>',
    ]
    if title:
        parts.append(
            f'<text x="{svg_width / 2:.1f}" y="28" text-anchor="middle" font-family="Arial, sans-serif" font-size="20" fill="#1F2937">{html.escape(title)}</text>'
        )

    for i in range(6):
        ratio = i / 5
        y = margin_top + ratio * plot_height
        value = y_max - ratio * (y_max - y_min)
        parts.append(f'<line x1="{margin_left}" y1="{y:.1f}" x2="{margin_left + plot_width}" y2="{y:.1f}" stroke="#D8DEE9" stroke-width="1"/>')
        parts.append(f'<text x="{margin_left - 10}" y="{y + 4:.1f}" text-anchor="end" font-family="Arial, sans-serif" font-size="12" fill="#4B5563">{_format_tick(value)}</text>')

    for i in range(6):
        ratio = i / 5
        x = margin_left + ratio * plot_width
        value = x_min + ratio * (x_max - x_min)
        parts.append(f'<line x1="{x:.1f}" y1="{margin_top}" x2="{x:.1f}" y2="{margin_top + plot_height}" stroke="#EEF2F7" stroke-width="1"/>')
        parts.append(f'<text x="{x:.1f}" y="{margin_top + plot_height + 22}" text-anchor="middle" font-family="Arial, sans-serif" font-size="12" fill="#4B5563">{_format_tick(value)}</text>')

    parts.extend([
        f'<line x1="{margin_left}" y1="{margin_top + plot_height}" x2="{margin_left + plot_width}" y2="{margin_top + plot_height}" stroke="#333333" stroke-width="1.4"/>',
        f'<line x1="{margin_left}" y1="{margin_top}" x2="{margin_left}" y2="{margin_top + plot_height}" stroke="#333333" stroke-width="1.4"/>',
        f'<text x="{margin_left + plot_width / 2:.1f}" y="{svg_height - 18}" text-anchor="middle" font-family="Arial, sans-serif" font-size="14" fill="#1F2937">{html.escape(x_label or x_name)}</text>',
        f'<text x="{margin_left}" y="{margin_top - 16}" text-anchor="start" font-family="Arial, sans-serif" font-size="14" fill="#1F2937">{html.escape(y_label or (series[0][0] if len(series) == 1 else "value"))}</text>',
    ])

    for idx, (name, values) in enumerate(series):
        color = MATLAB_COLORS[idx % len(MATLAB_COLORS)]
        points = [(sx(x), sy(y)) for x, y in zip(x_values, values)]
        if kind == "scatter":
            for x, y in points:
                parts.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="4.2" fill="{color}"/>')
        elif kind == "bar":
            slot_width = max(8, min(34, plot_width / max(len(x_values), 1) * 0.55))
            bar_width = slot_width / max(len(series), 1)
            zero_y = sy(max(0, min(y_max, max(y_min, 0))))
            for point_idx, (x_raw, y_raw) in enumerate(zip(x_values, values)):
                x = sx(x_raw) - slot_width / 2 + idx * bar_width
                y = sy(y_raw)
                top = min(y, zero_y)
                height_px = abs(zero_y - y)
                parts.append(f'<rect x="{x:.1f}" y="{top:.1f}" width="{bar_width - 1:.1f}" height="{height_px:.1f}" fill="{color}"/>')
        else:
            path = " ".join(
                f"{'M' if point_idx == 0 else 'L'} {x:.1f} {y:.1f}"
                for point_idx, (x, y) in enumerate(points)
            )
            parts.append(f'<path d="{path}" fill="none" stroke="{color}" stroke-width="2.6" stroke-linejoin="round" stroke-linecap="round"/>')
            for x, y in points:
                parts.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3.6" fill="#FFFFFF" stroke="{color}" stroke-width="2"/>')

    if len(series) > 1:
        legend_x = margin_left + plot_width - 130
        legend_y = margin_top + 14
        for idx, (name, _) in enumerate(series):
            y = legend_y + idx * 20
            color = MATLAB_COLORS[idx % len(MATLAB_COLORS)]
            parts.append(f'<line x1="{legend_x}" y1="{y}" x2="{legend_x + 18}" y2="{y}" stroke="{color}" stroke-width="3"/>')
            parts.append(f'<text x="{legend_x + 24}" y="{y + 4}" font-family="Arial, sans-serif" font-size="12" fill="#1F2937">{html.escape(name)}</text>')

    parts.append("</svg>")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(parts), encoding="utf-8")


def render_plot(
    input_path: Path,
    output_path: Path,
    *,
    kind: str,
    x_column: str | None,
    y_columns: list[str],
    title: str | None,
    x_label: str | None,
    y_label: str | None,
    width: float,
    height: float,
    dpi: int,
) -> None:
    rows = load_rows(input_path)
    if not rows:
        raise ValueError("Input data is empty")

    selected_y = _choose_y_columns(rows, y_columns)
    x_name, x_values = _choose_x_values(rows, x_column, selected_y)
    series: list[tuple[str, list[float]]] = []
    for column in selected_y:
        raw_values = _numeric_values(rows, column)
        if any(value is None for value in raw_values):
            raise ValueError(f"y column must be numeric: {column}")
        series.append((column, [float(value) for value in raw_values if value is not None]))

    output_format = output_path.suffix.lower().lstrip(".") or "svg"
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as exc:
        if output_format == "svg":
            _render_svg_fallback(
                output_path,
                kind=kind,
                x_name=x_name,
                x_values=x_values,
                series=series,
                title=title,
                x_label=x_label,
                y_label=y_label,
                width=width,
                height=height,
                dpi=dpi,
            )
            return
        raise RuntimeError("PNG output requires matplotlib; SVG output can run without it") from exc

    plt.rcParams.update({
        "font.family": ["Arial", "Microsoft YaHei", "SimHei", "sans-serif"],
        "axes.edgecolor": "#333333",
        "axes.grid": True,
        "grid.color": "#D8DEE9",
        "grid.linewidth": 0.8,
        "grid.alpha": 0.8,
        "lines.linewidth": 2.2,
        "figure.facecolor": "#FFFFFF",
        "axes.facecolor": "#FFFFFF",
    })

    fig, ax = plt.subplots(figsize=(width, height), dpi=dpi)
    for i, (column, y_values) in enumerate(series):
        color = MATLAB_COLORS[i % len(MATLAB_COLORS)]
        if kind == "scatter":
            ax.scatter(x_values, y_values, color=color, label=column, s=28)
        elif kind == "bar":
            offset = (i - (len(selected_y) - 1) / 2) * 0.22
            ax.bar([x + offset for x in x_values], y_values, width=0.2, color=color, label=column)
        else:
            ax.plot(x_values, y_values, color=color, marker="o", markersize=4, label=column)

    if title:
        ax.set_title(title, fontsize=14, pad=12)
    ax.set_xlabel(x_label or x_name)
    ax.set_ylabel(y_label or (selected_y[0] if len(selected_y) == 1 else "value"))
    if len(selected_y) > 1:
        ax.legend(frameon=False)
    fig.tight_layout()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, format=output_format, dpi=dpi)
    plt.close(fig)


def main() -> int:
    parser = argparse.ArgumentParser(description="Render MATLAB/Matplotlib-style plot assets.")
    parser.add_argument("input", type=Path, help="CSV, JSON, XLS, or XLSX data file")
    parser.add_argument("--output", "-o", type=Path, help="Output SVG/PNG path")
    parser.add_argument("--kind", choices=["line", "scatter", "bar"], default="line")
    parser.add_argument("--x", dest="x_column", help="Numeric x-axis column")
    parser.add_argument("--y", dest="y_columns", action="append", default=[], help="Numeric y-axis column; repeatable")
    parser.add_argument("--title")
    parser.add_argument("--x-label")
    parser.add_argument("--y-label")
    parser.add_argument("--width", type=float, default=7.2)
    parser.add_argument("--height", type=float, default=4.2)
    parser.add_argument("--dpi", type=int, default=160)
    args = parser.parse_args()

    output = args.output or args.input.with_name(f"{args.input.stem}_plot.svg")
    render_plot(
        args.input,
        output,
        kind=args.kind,
        x_column=args.x_column,
        y_columns=args.y_columns,
        title=args.title,
        x_label=args.x_label,
        y_label=args.y_label,
        width=args.width,
        height=args.height,
        dpi=args.dpi,
    )
    print(f"plot asset written: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
