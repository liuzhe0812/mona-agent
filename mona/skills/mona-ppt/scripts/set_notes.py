# -*- coding: utf-8 -*-
"""
批量为 PPTX 演示文稿的每一页设置演讲备注（Speaker Notes）的快捷脚本。

典型场景：LLM Agent 收到"为每页 PPT 加 100 字左右演讲备注"的需求时，
只需生成一份 notes.json（1-based 幻灯片编号 -> 备注文本），
再调用一次本脚本即可，无需手写 XML 或自行编写 python-pptx 代码。

用法：
    python scripts/set_notes.py <input.pptx> <output.pptx> --notes-file <notes.json>
    python scripts/set_notes.py <input.pptx> <output.pptx> --notes-inline '{"1":"..."}'
    python scripts/set_notes.py <input.pptx> <output.pptx> --notes-txt <notes.txt>

notes.json 示例（UTF-8 编码）：
    {
        "1": "第 1 页的演讲备注。\\n可以使用 \\n 表示换行，会在备注中呈现为多段落。",
        "2": "第 2 页的备注...",
        "3": ""
    }

notes.txt 示例（UTF-8 编码，推荐！彻底避免 JSON 引号转义陷阱）：
    ===SLIDE 1===
    第 1 页的演讲备注。
    可以直接换行，无需转义。中文全角引号"原样写入"也没问题。
    ===SLIDE 2===
    第 2 页的备注...
    ===SLIDE 3===

说明：
- 分隔符行格式严格为 `===SLIDE N===`（N 为 1-based 整数），其下一行起到下一个分隔符之前的全部文本即为该页备注。
- 首个分隔符之前的任何内容都会被忽略（可用于写说明/备忘）。
- 每段末尾的多余空行会被自动去除。某段无内容（如示例中的第 3 页）= 清空该页备注。

行为说明：
- 键为 1-based 的幻灯片编号，允许字符串或整数；非数字键（如 "cover"）会被跳过并告警。
- 值为纯文本；空字符串会清空该页备注；null 表示跳过该页。
- 已有旧备注会被**完全覆盖**（整篇替换，不做追加）。
- 原本没有 notesSlide 的幻灯片会在写入时自动创建（由 python-pptx 负责）。
- output.pptx 必须与 input.pptx 不同路径，避免覆盖原文件。

依赖：仅 Python 标准库 + python-pptx（pip install python-pptx）。
"""

import argparse
import json
import re
import sys
from pathlib import Path

from pptx import Presentation


PREVIEW_LIMIT = 30


def _build_parser() -> argparse.ArgumentParser:
    """构造命令行参数解析器。"""
    parser = argparse.ArgumentParser(
        prog="set_notes.py",
        description="批量为 PPTX 每一页设置演讲备注（Speaker Notes）。",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "notes.json 示例：\n"
            '  {\n'
            '      "1": "第 1 页备注文本。\\n支持 \\\\n 换行",\n'
            '      "2": "第 2 页备注文本"\n'
            '  }\n'
            "说明：键为 1-based 幻灯片编号（字符串或整数均可），值为纯文本备注。"
        ),
    )
    parser.add_argument("input_pptx", help="输入 .pptx 文件路径")
    parser.add_argument("output_pptx", help="输出 .pptx 文件路径（不能与输入相同）")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--notes-file", help="包含备注映射的 JSON 文件路径（UTF-8）")
    group.add_argument("--notes-inline", help="直接传入 JSON 字符串（便于短文本快速调用）")
    group.add_argument(
        "--notes-txt",
        help=(
            "纯文本备注文件路径（UTF-8），以 `===SLIDE N===` 行作为分隔符。"
            "推荐用于大段中文备注，彻底避免 JSON 引号转义陷阱。"
        ),
    )
    parser.add_argument(
        "--quiet",
        action="store_true",
        help="抑制逐页预览输出，仅打印汇总行",
    )
    return parser


_SLIDE_HEADER_RE = re.compile(r"^\s*={3,}\s*SLIDE\s+(\d+)\s*={3,}\s*$", re.IGNORECASE)


def _parse_notes_txt(text: str) -> dict:
    """解析 `===SLIDE N===` 分隔符格式的纯文本，返回 {编号字符串: 备注文本}。"""
    result: dict = {}
    current_idx: int | None = None
    buffer: list = []

    def _flush() -> None:
        if current_idx is None:
            return
        # 去除每段末尾多余的空行，但保留段落内部的换行结构
        content = "\n".join(buffer).rstrip("\n")
        result[str(current_idx)] = content

    for line in text.splitlines():
        match = _SLIDE_HEADER_RE.match(line)
        if match:
            _flush()
            current_idx = int(match.group(1))
            buffer = []
        else:
            if current_idx is not None:
                buffer.append(line)
    _flush()

    if not result:
        print(
            "错误：--notes-txt 未解析到任何 `===SLIDE N===` 分隔符段落",
            file=sys.stderr,
        )
        sys.exit(1)
    return result


def _load_notes(args: argparse.Namespace) -> dict:
    """根据参数加载并返回 notes 映射（JSON 解析后的原始 dict）。"""
    if args.notes_file:
        notes_path = Path(args.notes_file)
        if not notes_path.is_file():
            print(f"错误：--notes-file 指定的文件不存在：{notes_path}", file=sys.stderr)
            sys.exit(1)
        try:
            with notes_path.open("r", encoding="utf-8") as handle:
                data = json.load(handle)
        except json.JSONDecodeError as exc:
            print(f"错误：notes JSON 解析失败：{exc}", file=sys.stderr)
            sys.exit(1)
    elif args.notes_txt:
        txt_path = Path(args.notes_txt)
        if not txt_path.is_file():
            print(f"错误：--notes-txt 指定的文件不存在：{txt_path}", file=sys.stderr)
            sys.exit(1)
        try:
            with txt_path.open("r", encoding="utf-8") as handle:
                raw_text = handle.read()
        except OSError as exc:
            print(f"错误：读取 notes-txt 文件失败：{exc}", file=sys.stderr)
            sys.exit(1)
        data = _parse_notes_txt(raw_text)
    else:
        try:
            data = json.loads(args.notes_inline)
        except json.JSONDecodeError as exc:
            print(f"错误：--notes-inline JSON 解析失败：{exc}", file=sys.stderr)
            sys.exit(1)

    if not isinstance(data, dict):
        print("错误：notes 内容必须是 JSON 对象（{编号: 文本}）", file=sys.stderr)
        sys.exit(1)
    return data


def _validate_paths(input_pptx: Path, output_pptx: Path) -> None:
    """校验输入输出路径的合法性。"""
    if not input_pptx.is_file():
        print(f"错误：输入 pptx 文件不存在：{input_pptx}", file=sys.stderr)
        sys.exit(1)
    if input_pptx.resolve() == output_pptx.resolve():
        print("错误：output 路径不能与 input 相同，禁止覆盖原文件", file=sys.stderr)
        sys.exit(1)


def _apply_notes(prs: Presentation, notes: dict, quiet: bool) -> int:
    """把 notes 映射写入到 prs 的对应幻灯片中，返回成功写入的页数。"""
    total_slides = len(prs.slides)
    applied = 0
    # 对键进行稳定排序（数字键按数值顺序），便于预览输出可读
    items = []
    for raw_key, text in notes.items():
        try:
            idx = int(raw_key)
        except (TypeError, ValueError):
            print(f"警告：忽略非数字键 {raw_key!r}", file=sys.stderr)
            continue
        items.append((idx, text))
    items.sort(key=lambda pair: pair[0])

    for idx, text in items:
        if text is None:
            continue
        if idx < 1 or idx > total_slides:
            print(
                f"警告：幻灯片编号 {idx} 超出范围（1-{total_slides}），已跳过",
                file=sys.stderr,
            )
            continue
        if not isinstance(text, str):
            print(f"警告：幻灯片 {idx} 的备注不是字符串类型，已跳过", file=sys.stderr)
            continue

        slide = prs.slides[idx - 1]
        # 访问 notes_slide 属性时，python-pptx 会在缺失时自动创建 notesSlide
        slide.notes_slide.notes_text_frame.text = text
        applied += 1

        if not quiet:
            preview = text.replace("\n", " ")
            if len(preview) > PREVIEW_LIMIT:
                preview = preview[:PREVIEW_LIMIT] + "..."
            print(f"slide {idx}: {preview}")

    return applied


def main() -> int:
    """脚本主入口。"""
    parser = _build_parser()
    args = parser.parse_args()

    input_pptx = Path(args.input_pptx)
    output_pptx = Path(args.output_pptx)

    _validate_paths(input_pptx, output_pptx)
    notes = _load_notes(args)

    try:
        prs = Presentation(str(input_pptx))
        applied = _apply_notes(prs, notes, args.quiet)
        output_pptx.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_pptx))
    except (OSError, ValueError, KeyError) as exc:
        print(f"错误：处理 pptx 失败：{exc}", file=sys.stderr)
        return 1
    except Exception as exc:  # pylint: disable=broad-except
        # python-pptx / lxml 可能抛出形态各异的异常，这里兜底，确保中文错误输出
        print(f"错误：未预期的异常：{exc}", file=sys.stderr)
        return 1

    print(f"已为 {applied} 张幻灯片设置备注，输出：{output_pptx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
