#!/usr/bin/env python3
"""Build a PPTX from a template and content plan (Native PPTX Template Mode).

Usage:
    python native_pptx_builder.py \
      --template-dir <template_dir> \
      --plan <native_content_plan.json> \
      --output <output.pptx>
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
SLIDE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
)
SLIDE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
)

ET.register_namespace("", PRESENTATION_NS)
ET.register_namespace("r", REL_NS)


def _zone_to_emu(
    zone: dict[str, float], slide_width: int, slide_height: int
) -> tuple[int, int, int, int]:
    """Convert normalized zone coordinates to EMU."""
    return (
        int(zone["x"] * slide_width),
        int(zone["y"] * slide_height),
        int(zone["w"] * slide_width),
        int(zone["h"] * slide_height),
    )


def _add_text_box(
    slide, text: str, left: int, top: int, width: int, height: int, font_size: int = 18
) -> None:
    """Add a text box to a slide."""
    from pptx.util import Pt

    txbox = slide.shapes.add_textbox(left, top, width, height)
    tf = txbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    for run in p.runs:
        run.font.size = Pt(font_size)


def _add_bullets(
    slide, items: list[str], left: int, top: int, width: int, height: int, font_size: int = 14
) -> None:
    """Add a text box with bullet points to a slide."""
    from pptx.util import Pt

    txbox = slide.shapes.add_textbox(left, top, width, height)
    tf = txbox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.level = 0
        for run in p.runs:
            run.font.size = Pt(font_size)


def _add_image(slide, image_path: str, left: int, top: int, width: int, height: int) -> None:
    """Add an image to a slide."""
    p = Path(image_path)
    if p.exists():
        slide.shapes.add_picture(str(p), left, top, width, height)


def _resolve_image_path(image_path: str, plan_path: Path, template_dir: Path) -> Path | None:
    """Resolve image paths from common project locations."""
    p = Path(image_path)
    candidates = [p] if p.is_absolute() else [
        plan_path.parent / p,
        plan_path.parent / "images" / p,
        template_dir / p,
        Path.cwd() / p,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def _add_image_if_exists(
    slide,
    image_path: str,
    plan_path: Path,
    template_dir: Path,
    left: int,
    top: int,
    width: int,
    height: int,
) -> None:
    p = _resolve_image_path(image_path, plan_path, template_dir)
    if p is not None:
        slide.shapes.add_picture(str(p), left, top, width, height)


def _slide_part_name(index: int) -> str:
    return f"ppt/slides/slide{index}.xml"


def _slide_rels_name(index: int) -> str:
    return f"ppt/slides/_rels/slide{index}.xml.rels"


def _read_required(zf: zipfile.ZipFile, name: str) -> bytes:
    try:
        return zf.read(name)
    except KeyError as exc:
        raise ValueError(f"template is missing {name}") from exc


def _max_numeric_rid(relationships_root: ET.Element) -> int:
    max_id = 0
    for rel in relationships_root:
        rid = rel.attrib.get("Id", "")
        m = re.fullmatch(r"rId(\d+)", rid)
        if m:
            max_id = max(max_id, int(m.group(1)))
    return max_id


def _rewrite_content_types(xml: bytes, slide_count: int) -> bytes:
    root = ET.fromstring(xml)
    override_tag = f"{{{CONTENT_TYPES_NS}}}Override"
    for child in list(root):
        if (
            child.tag == override_tag
            and re.fullmatch(r"/ppt/slides/slide\d+\.xml", child.attrib.get("PartName", ""))
        ):
            root.remove(child)
    for index in range(1, slide_count + 1):
        ET.SubElement(
            root,
            override_tag,
            {
                "PartName": f"/ppt/slides/slide{index}.xml",
                "ContentType": SLIDE_CONTENT_TYPE,
            },
        )
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _rewrite_presentation(xml: bytes, slide_rids: list[str]) -> bytes:
    root = ET.fromstring(xml)
    sld_id_lst = root.find(f"{{{PRESENTATION_NS}}}sldIdLst")
    if sld_id_lst is None:
        sld_id_lst = ET.SubElement(root, f"{{{PRESENTATION_NS}}}sldIdLst")
    for child in list(sld_id_lst):
        sld_id_lst.remove(child)
    for idx, rid in enumerate(slide_rids, start=1):
        ET.SubElement(
            sld_id_lst,
            f"{{{PRESENTATION_NS}}}sldId",
            {"id": str(255 + idx), f"{{{REL_NS}}}id": rid},
        )
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _rewrite_presentation_rels(xml: bytes, slide_count: int) -> tuple[bytes, list[str]]:
    root = ET.fromstring(xml)
    rel_tag = f"{{{PKG_REL_NS}}}Relationship"
    for child in list(root):
        if child.attrib.get("Type") == SLIDE_REL_TYPE:
            root.remove(child)

    start = _max_numeric_rid(root) + 1
    slide_rids = [f"rId{start + i}" for i in range(slide_count)]
    for idx, rid in enumerate(slide_rids, start=1):
        ET.SubElement(
            root,
            rel_tag,
            {
                "Id": rid,
                "Type": SLIDE_REL_TYPE,
                "Target": f"slides/slide{idx}.xml",
            },
        )
    return ET.tostring(root, encoding="utf-8", xml_declaration=True), slide_rids


def _clone_template_slides(
    template_pptx: Path,
    role_sequence: list[str],
    role_to_slide: dict[str, int],
    output_path: Path,
) -> None:
    """Create a PPTX whose slide order follows roles from the content plan."""
    if not role_sequence:
        raise ValueError("role sequence is empty")

    with zipfile.ZipFile(template_pptx, "r") as zin:
        original_names = set(zin.namelist())
        slide_xml_by_source: dict[int, bytes] = {}
        slide_rels_by_source: dict[int, bytes | None] = {}

        source_indices: list[int] = []
        for role in role_sequence:
            source_idx = role_to_slide.get(role) or role_to_slide.get("content") or 1
            source_indices.append(source_idx)
            if source_idx not in slide_xml_by_source:
                slide_xml_by_source[source_idx] = _read_required(
                    zin, _slide_part_name(source_idx)
                )
                rels_name = _slide_rels_name(source_idx)
                slide_rels_by_source[source_idx] = (
                    zin.read(rels_name) if rels_name in original_names else None
                )

        presentation_rels_xml, slide_rids = _rewrite_presentation_rels(
            _read_required(zin, "ppt/_rels/presentation.xml.rels"),
            len(role_sequence),
        )
        presentation_xml = _rewrite_presentation(
            _read_required(zin, "ppt/presentation.xml"),
            slide_rids,
        )
        content_types_xml = _rewrite_content_types(
            _read_required(zin, "[Content_Types].xml"),
            len(role_sequence),
        )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                name = item.filename
                if name == "[Content_Types].xml":
                    zout.writestr(item, content_types_xml)
                elif name == "ppt/presentation.xml":
                    zout.writestr(item, presentation_xml)
                elif name == "ppt/_rels/presentation.xml.rels":
                    zout.writestr(item, presentation_rels_xml)
                elif re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                    continue
                elif re.fullmatch(r"ppt/slides/_rels/slide\d+\.xml\.rels", name):
                    continue
                else:
                    zout.writestr(item, zin.read(name))

            for dst_idx, src_idx in enumerate(source_indices, start=1):
                zout.writestr(_slide_part_name(dst_idx), slide_xml_by_source[src_idx])
                rels_xml = slide_rels_by_source[src_idx]
                if rels_xml is not None:
                    zout.writestr(_slide_rels_name(dst_idx), rels_xml)


def build_pptx(template_dir: Path, plan_path: Path, output_path: Path) -> int:
    """Build PPTX from template and content plan."""
    from pptx import Presentation

    # Load template data
    template_pptx = template_dir / "template.pptx"
    roles_path = template_dir / "template_roles.json"
    manifest_path = template_dir / "template_manifest.json"

    if not template_pptx.exists():
        print(f"Error: template.pptx not found in {template_dir}", file=sys.stderr)
        return 1
    if not roles_path.exists():
        print(f"Error: template_roles.json not found in {template_dir}", file=sys.stderr)
        return 1

    roles_data = json.loads(roles_path.read_text(encoding="utf-8"))
    role_to_slide = {
        str(role): int(index)
        for role, index in roles_data.get("roles", {}).items()
        if isinstance(index, int) and index > 0
    }
    zones = roles_data.get("zones", {})

    manifest_data = {}
    if manifest_path.exists():
        manifest_data = json.loads(manifest_path.read_text(encoding="utf-8"))

    slide_width = manifest_data.get("slide_width_emu", 12192000)
    slide_height = manifest_data.get("slide_height_emu", 6858000)

    # Load content plan
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    plan_slides = plan.get("slides", [])

    if not plan_slides:
        print("Error: content plan has no slides", file=sys.stderr)
        return 1

    # Build the target slide sequence
    target_slides: list[tuple[dict, str]] = []  # (plan_slide, role)
    for ps in plan_slides:
        role = ps.get("role", "content")
        if role not in role_to_slide:
            role = "content" if "content" in role_to_slide else "cover"
        target_slides.append((ps, role))

    with tempfile.TemporaryDirectory(prefix="native_pptx_build_") as tmp_dir:
        cloned_path = Path(tmp_dir) / "cloned.pptx"
        try:
            _clone_template_slides(
                template_pptx,
                [role for _, role in target_slides],
                role_to_slide or {"content": 1},
                cloned_path,
            )
        except ValueError as exc:
            print(f"Error: {exc}", file=sys.stderr)
            return 1
        prs = Presentation(str(cloned_path))

        # Fill content on cloned slides
        for i, (ps, role) in enumerate(target_slides):
            slide = prs.slides[i]
            zone = zones.get(role, zones.get("content", {}))

            # Fill title
            title_zone = zone.get("title")
            if title_zone and ps.get("title"):
                left, top, width, height = _zone_to_emu(title_zone, slide_width, slide_height)
                _add_text_box(slide, ps["title"], left, top, width, height, font_size=24)

            # Fill subtitle (cover/thanks)
            subtitle_zone = zone.get("subtitle")
            if subtitle_zone and ps.get("subtitle"):
                left, top, width, height = _zone_to_emu(
                    subtitle_zone, slide_width, slide_height
                )
                _add_text_box(slide, ps["subtitle"], left, top, width, height, font_size=16)

            # Fill body (toc items / content bullets)
            body_zone = zone.get("body")
            if body_zone:
                left, top, width, height = _zone_to_emu(body_zone, slide_width, slide_height)
                items = ps.get("items") or ps.get("bullets")
                if items:
                    _add_bullets(slide, items, left, top, width, height, font_size=14)

                # Fill image if present
                image_path = ps.get("image")
                if image_path:
                    # Scale image to fit in body zone (use 60% of zone height for image)
                    img_height = int(height * 0.6)
                    img_top = top + height - img_height
                    _add_image_if_exists(
                        slide,
                        image_path,
                        plan_path,
                        template_dir,
                        left,
                        img_top,
                        width,
                        img_height,
                    )

        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))

    print(f"[OK] Native PPTX built: {output_path}")
    print(f"     Plan slides: {len(target_slides)}")
    print(f"     Output slides: {len(prs.slides)}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Build PPTX from template and content plan")
    parser.add_argument("--template-dir", required=True, help="Native template directory")
    parser.add_argument("--plan", required=True, help="Path to native_content_plan.json")
    parser.add_argument("--output", required=True, help="Output PPTX file path")
    args = parser.parse_args()

    template_dir = Path(args.template_dir)
    plan_path = Path(args.plan)
    output_path = Path(args.output)

    if not template_dir.is_dir():
        print(f"Error: template directory not found: {template_dir}", file=sys.stderr)
        return 1
    if not plan_path.exists():
        print(f"Error: plan file not found: {plan_path}", file=sys.stderr)
        return 1

    return build_pptx(template_dir, plan_path, output_path)


if __name__ == "__main__":
    raise SystemExit(main())
