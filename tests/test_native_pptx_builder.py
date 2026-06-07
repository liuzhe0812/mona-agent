from __future__ import annotations

import importlib.util
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
BUILDER_PATH = ROOT / "mona" / "skills" / "mona-ppt" / "scripts" / "native_pptx_builder.py"


def _load_builder():
    spec = importlib.util.spec_from_file_location("native_pptx_builder", BUILDER_PATH)
    assert spec is not None
    assert spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _slide_text(slide) -> str:
    chunks: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            chunks.append(shape.text)
    return "\n".join(chunks)


def test_native_builder_repeats_content_role_slide(tmp_path: Path) -> None:
    template_dir = tmp_path / "template"
    template_dir.mkdir()

    prs = Presentation()
    for index, role in enumerate(["cover", "toc", "content", "thanks"], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        marker = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(4), Inches(0.5))
        marker.text_frame.text = f"template {index} {role}"
    prs.save(template_dir / "template.pptx")

    roles = {
        "roles": {"cover": 1, "toc": 2, "content": 3, "thanks": 4},
        "zones": {
            "cover": {
                "title": {"x": 0.1, "y": 0.1, "w": 0.8, "h": 0.12},
                "subtitle": {"x": 0.1, "y": 0.25, "w": 0.8, "h": 0.1},
            },
            "toc": {
                "title": {"x": 0.1, "y": 0.1, "w": 0.8, "h": 0.12},
                "body": {"x": 0.1, "y": 0.25, "w": 0.8, "h": 0.5},
            },
            "content": {
                "title": {"x": 0.1, "y": 0.1, "w": 0.8, "h": 0.12},
                "body": {"x": 0.1, "y": 0.25, "w": 0.8, "h": 0.5},
            },
            "thanks": {
                "title": {"x": 0.1, "y": 0.1, "w": 0.8, "h": 0.12},
                "subtitle": {"x": 0.1, "y": 0.25, "w": 0.8, "h": 0.1},
            },
        },
    }
    (template_dir / "template_roles.json").write_text(
        json.dumps(roles), encoding="utf-8"
    )
    (template_dir / "template_manifest.json").write_text(
        json.dumps({"slide_width_emu": 12192000, "slide_height_emu": 6858000}),
        encoding="utf-8",
    )

    plan = {
        "title": "native test",
        "slides": [
            {"role": "cover", "title": "Cover", "subtitle": "Sub"},
            {"role": "toc", "title": "TOC", "items": ["A", "B"]},
            {"role": "content", "title": "Content 1", "bullets": ["One"]},
            {"role": "content", "title": "Content 2", "bullets": ["Two"]},
            {"role": "content", "title": "Content 3", "bullets": ["Three"]},
            {"role": "thanks", "title": "Thanks", "subtitle": "End"},
        ],
    }
    plan_path = tmp_path / "native_content_plan.json"
    plan_path.write_text(json.dumps(plan), encoding="utf-8")
    output_path = tmp_path / "out.pptx"

    builder = _load_builder()
    assert builder.build_pptx(template_dir, plan_path, output_path) == 0

    out = Presentation(str(output_path))
    assert len(out.slides) == 6
    assert "template 1 cover" in _slide_text(out.slides[0])
    assert "template 2 toc" in _slide_text(out.slides[1])
    assert "template 3 content" in _slide_text(out.slides[2])
    assert "template 3 content" in _slide_text(out.slides[3])
    assert "template 3 content" in _slide_text(out.slides[4])
    assert "template 4 thanks" in _slide_text(out.slides[5])
