"""Tests for Office media-to-source context mapping."""

from __future__ import annotations

import base64
import io
import zipfile
from pathlib import Path

from mona.materials.visual_context import office_image_contexts

_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def test_docx_table_image_keeps_same_row_text(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "软终端/H5终端"
    table.cell(0, 1).text = "通过现有电脑系统软终端或浏览器html5访问"
    table.cell(0, 2).paragraphs[0].add_run().add_picture(io.BytesIO(_PNG_1X1))
    path = tmp_path / "table.docx"
    document.save(path)

    contexts = office_image_contexts(path)
    media = next(key for key in contexts if key.startswith("word/media/"))
    assert "软终端/H5终端" in contexts[media]
    assert "通过现有电脑系统软终端或浏览器html5访问" in contexts[media]


def test_docx_vml_image_gets_neighboring_paragraphs(tmp_path: Path) -> None:
    document_xml = """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:v="urn:schemas-microsoft-com:vml"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
 <w:body>
  <w:p><w:r><w:t>前一段说明</w:t></w:r></w:p>
  <w:p><w:r><w:pict><v:shape><v:imagedata r:id="rId1"/></v:shape></w:pict></w:r></w:p>
  <w:p><w:r><w:t>后一段说明</w:t></w:r></w:p>
 </w:body>
</w:document>"""
    rels_xml = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
 <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
  Target="media/image1.png"/>
</Relationships>"""
    path = tmp_path / "vml.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/_rels/document.xml.rels", rels_xml)
        archive.writestr("word/media/image1.png", _PNG_1X1)

    contexts = office_image_contexts(path)
    assert "前一段说明" in contexts["word/media/image1.png"]
    assert "后一段说明" in contexts["word/media/image1.png"]


def test_pptx_image_includes_slide_text(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "软终端方案"
    slide.shapes.add_picture(io.BytesIO(_PNG_1X1), Inches(1), Inches(2))
    path = tmp_path / "slides.pptx"
    presentation.save(path)

    contexts = office_image_contexts(path)
    context = contexts["ppt/media/image1.png"]
    assert "第 1 张幻灯片" in context
    assert "软终端方案" in context


def test_pptx_uses_presentation_order_for_slide_number(tmp_path: Path) -> None:
    from xml.etree import ElementTree

    from PIL import Image as PillowImage
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    second_image = io.BytesIO()
    PillowImage.new("RGB", (1, 1), "blue").save(second_image, "PNG")
    for label, image_data in (("第一张", _PNG_1X1), ("第二张", second_image.getvalue())):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        textbox.text = label
        slide.shapes.add_picture(io.BytesIO(image_data), Inches(1), Inches(2))
    original = tmp_path / "ordered.pptx"
    presentation.save(original)

    with zipfile.ZipFile(original) as source:
        members = {name: source.read(name) for name in source.namelist()}
    root = ElementTree.fromstring(members["ppt/presentation.xml"])
    slide_id_list = next(element for element in root.iter() if element.tag.endswith("}sldIdLst"))
    slide_id_list[:] = list(reversed(list(slide_id_list)))
    members["ppt/presentation.xml"] = ElementTree.tostring(
        root, encoding="utf-8", xml_declaration=True,
    )
    reordered = tmp_path / "reordered.pptx"
    with zipfile.ZipFile(reordered, "w") as target:
        for name, data in members.items():
            target.writestr(name, data)

    contexts = office_image_contexts(reordered)
    assert "第 2 张幻灯片" in contexts["ppt/media/image1.png"]
    assert "第一张" in contexts["ppt/media/image1.png"]
    assert "第 1 张幻灯片" in contexts["ppt/media/image2.png"]
    assert "第二张" in contexts["ppt/media/image2.png"]


def test_xlsx_image_includes_anchor_and_nearby_values(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "终端配置"
    worksheet["A1"] = "类型"
    worksheet["B2"] = "软终端"
    worksheet["C3"] = "浏览器html5"
    worksheet.add_image(Image(io.BytesIO(_PNG_1X1)), "B2")
    path = tmp_path / "sheet.xlsx"
    workbook.save(path)
    workbook.close()

    contexts = office_image_contexts(path)
    context = contexts["xl/media/image1.png"]
    assert "终端配置" in context
    assert "B2" in context
    assert "软终端" in context
    assert "浏览器html5" in context


def test_xlsx_uses_drawing_anchor_when_openpyxl_drops_an_image(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "终端配置"
    worksheet["B2"] = "第一张图片"
    worksheet["D4"] = "第二张图片"
    worksheet.add_image(Image(io.BytesIO(_PNG_1X1)), "B2")
    worksheet.add_image(Image(io.BytesIO(_PNG_1X1)), "D4")
    original = tmp_path / "images.xlsx"
    workbook.save(original)
    workbook.close()

    with zipfile.ZipFile(original) as source:
        members = {name: source.read(name) for name in source.namelist()}
    members["xl/media/image2.png"] = b"not a decodable image"
    broken = tmp_path / "broken-image.xlsx"
    with zipfile.ZipFile(broken, "w") as target:
        for name, data in members.items():
            target.writestr(name, data)

    contexts = office_image_contexts(broken)
    assert "D4" in contexts["xl/media/image2.png"]
    assert "第二张图片" in contexts["xl/media/image2.png"]


def test_context_is_bounded_and_marks_truncation(tmp_path: Path) -> None:
    long_text = "很长的上下文" * 700
    document_xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
 <w:body><w:p><w:r><w:t>{long_text}</w:t><w:drawing><a:blip r:embed="rId1"/></w:drawing></w:r></w:p></w:body>
</w:document>"""
    rels_xml = """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
 <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>
</Relationships>"""
    path = tmp_path / "long.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/_rels/document.xml.rels", rels_xml)
        archive.writestr("word/media/image1.png", _PNG_1X1)

    context = office_image_contexts(path)["word/media/image1.png"]
    assert len(context) <= 3000
    assert "截断" in context
