"""Contract checks for the professional Office specialist skills."""

from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import httpx
import pytest
from pydantic import ValidationError

from mona.agent.partners import MONA_AGENT_ID
from mona.agent.runner import AgentRunner, AgentRunSpec
from mona.agent.skills import SkillsLoader
from mona.agent.tools.context import RequestContext
from mona.agent.tools.office import OfficeTool
from mona.agent.tools.registry import ToolRegistry
from mona.office.api import OfficeSocketHub
from mona.office.client import OfficeServiceClient
from mona.office.schemas import (
    DocumentVersion,
    OfficeApplyCommand,
    OfficeCommandFailure,
    OfficeCommandSuccess,
    OfficeExportRequest,
    OfficeInspectRequest,
    OfficeInspectResult,
    OfficeInspectSuccess,
    OfficeSaveRequest,
    OfficeSessionCreateRequest,
    VisualResult,
)
from mona.providers.base import LLMProvider, LLMResponse, ToolCallRequest

SKILL_ROOT = Path(__file__).parents[2] / "mona" / "skills"
SPECIALIST_SKILLS = ("mona-docx", "mona-xlsx", "mona-pptx")
JSON_FENCE = re.compile(r"```json\s*\r?\n(.*?)\r?\n```", re.IGNORECASE | re.DOTALL)
OFFICE_ACTIONS = {"open", "inspect", "apply", "save", "export", "close"}


@pytest.mark.asyncio
async def test_word_advanced_helpers_roundtrip_via_skill_tool(tmp_path, monkeypatch):
    from docx import Document

    from mona.agent.tools import skill_tools
    from mona.runtime.agent_env import AgentEnvironmentResolution

    loader = SkillsLoader(tmp_path, builtin_skills_dir=SKILL_ROOT)
    monkeypatch.setattr(skill_tools, "_skills_loader", lambda _agent_id: loader)

    class LocalTestRuntime:
        async def prepare_for_skill(self, suffix, skill_dir, spec):
            assert suffix == ".py"
            assert skill_dir == SKILL_ROOT / "mona-docx"
            assert (skill_dir / "pyproject.toml").is_file()
            return AgentEnvironmentResolution(executable=Path(sys.executable), env={})

    tool = skill_tools.SkillScriptRunTool(
        track_usage=False, workspace=tmp_path, agent_environment=LocalTestRuntime(),
    )
    original = tmp_path / "原始文档.docx"
    unpacked = tmp_path / "解包"
    output = tmp_path / "修改结果.docx"
    document = Document()
    document.add_paragraph("Original content")
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Preserved table"
    document.save(original)
    source_bytes = original.read_bytes()

    result = await tool.execute(
        skill="mona-docx", script="unpack.py",
        args=f'"{original}" "{unpacked}" --merge-runs false --simplify-redlines false',
    )
    assert "[exit code: 0]" in result, result
    xml = unpacked / "word" / "document.xml"
    xml.write_text(
        xml.read_text(encoding="utf-8").replace("Original content", "Updated content"),
        encoding="utf-8",
    )
    result = await tool.execute(
        skill="mona-docx", script="pack.py",
        args=f'"{unpacked}" "{output}" --original "{original}"',
    )
    assert "[exit code: 0]" in result, result
    assert "All validations PASSED" in result, result
    restored = Document(output)
    assert restored.paragraphs[0].text == "Updated content"
    assert restored.tables[0].cell(0, 0).text == "Preserved table"
    assert original.read_bytes() == source_bytes
    with zipfile.ZipFile(output) as archive:
        assert archive.testzip() is None


@pytest.fixture
def builtin_loader_root() -> Path:
    return SKILL_ROOT


@pytest.fixture(autouse=True)
def isolate_agent_skill_paths(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from mona.config import paths as config_paths

    monkeypatch.setattr(
        config_paths,
        "get_agent_skills_dir",
        lambda _agent_id: tmp_path / "agent-skills",
    )


def _loader(
    tmp_path: Path,
    builtin_root: Path,
    agent_id: str,
) -> SkillsLoader:
    workspace = tmp_path / f"workspace-{agent_id.replace('.', '-')}"
    workspace.mkdir(exist_ok=True)
    return SkillsLoader(
        workspace=workspace,
        builtin_skills_dir=builtin_root,
        agent_id=agent_id,
    )


def test_real_loaders_discover_and_read_specialists_without_office_router(
    tmp_path: Path,
    builtin_loader_root: Path,
) -> None:
    named_agent = "com.mona.a-share-analyst"
    for agent_id in (MONA_AGENT_ID, named_agent):
        loader = _loader(tmp_path, builtin_loader_root, agent_id)
        names = {entry["name"] for entry in loader.list_skills(filter_unavailable=False)}
        assert set(SPECIALIST_SKILLS) <= names
        for name in SPECIALIST_SKILLS:
            content = loader.load_skill(name)
            metadata = loader.get_skill_metadata(name)
            assert content and content.startswith("---")
            assert metadata and metadata["name"] == name
        summary = loader.build_skills_summary()
        assert "prd-document" in summary
        for retired in ("docx", "mona-office", "doc-writing-guide"):
            assert retired not in names
            assert loader.resolve_skill_dir(retired) is None

    loader = _loader(tmp_path, builtin_loader_root, MONA_AGENT_ID)
    word_root = loader.resolve_skill_dir("mona-docx")
    assert word_root is not None
    assert (word_root / "references" / "advanced-ooxml.md").is_file()
    assert (word_root / "pyproject.toml").is_file()
    for script in ("unpack.py", "pack.py", "comment.py", "accept_changes.py"):
        assert (word_root / "scripts" / script).is_file()


def test_specialist_guidance_uses_incremental_reads_and_domain_acceptance() -> None:
    contents = {
        name: (SKILL_ROOT / name / "SKILL.md").read_text(encoding="utf-8")
        for name in SPECIALIST_SKILLS
    }

    for content in contents.values():
        assert "open" in content and "selection" in content and "version" in content
        assert "partial" in content and "覆盖清单" in content
        assert "重复读取" in content and "范围不重叠" in content
        assert "session_id" in content and "恢复同一会话" in content
        assert "inspect capabilities" in content and "operations: [" in content

    ppt = "\n".join(
        path.read_text(encoding="utf-8")
        for path in sorted((SKILL_ROOT / "mona-pptx").rglob("*.md"))
    )
    assert re.search(r"inspect\s+review", ppt)
    assert "pendingSlideIds" in ppt
    assert "REVIEW_REQUIRED" in ppt
    assert re.search(r"allow_unreviewed.*明确.*草稿", ppt)
    assert "自动 warning 也不全部是 bug" in ppt
    assert "operations: [" in ppt
    assert "使用顺序" in ppt and "不等待整套图片" in ppt
    assert "不要求把任意表格都改成原生对象" in ppt

    assert "REVIEW_REQUIRED" not in contents["mona-docx"]
    assert "REVIEW_REQUIRED" not in contents["mona-xlsx"]
    assert "原生单元格/范围用于需要编辑和数据含义" in contents["mona-xlsx"]
    assert "不要求任意表格都变成原生对象" in contents["mona-docx"]


def _office_json_examples() -> list[tuple[Path, dict[str, object]]]:
    examples: list[tuple[Path, dict[str, object]]] = []
    for skill_name in SPECIALIST_SKILLS:
        skill_dir = SKILL_ROOT / skill_name
        for path in sorted(skill_dir.rglob("*.md")):
            content = path.read_text(encoding="utf-8")
            for match in JSON_FENCE.finditer(content):
                try:
                    payload = json.loads(match.group(1))
                except json.JSONDecodeError as exc:
                    pytest.fail(f"invalid JSON example in {path}: {exc}")
                if isinstance(payload, dict) and payload.get("action") in OFFICE_ACTIONS:
                    examples.append((path, payload))
    return examples


def test_all_specialist_office_examples_validate_against_wire_models() -> None:
    examples = _office_json_examples()
    assert examples
    seen_skills: set[str] = set()
    for path, payload in examples:
        seen_skills.add(path.parent.parent.name if path.parent.name == "references" else path.parent.name)
        action = payload["action"]
        if action == "open":
            OfficeSessionCreateRequest.model_validate(
                {
                    "owner_session_key": "test-owner",
                    "type": payload.get("document_type"),
                    "path": payload.get("path"),
                    "display_name": payload.get("display_name"),
                }
            )
        elif action == "inspect":
            OfficeInspectRequest.model_validate(
                {
                    "session_id": payload["session_id"],
                    "query": payload["query"],
                }
            )
        elif action == "apply":
            OfficeApplyCommand.model_validate(
                {
                    "session_id": payload["session_id"],
                    "operation_id": "example-operation",
                    "expected_version": payload["expected_version"],
                    "operations": payload["operations"],
                }
            )
        elif action == "save":
            OfficeSaveRequest.model_validate(
                {
                    "version": payload.get("expected_version"),
                    "overwrite_source": payload.get("overwrite_source", False),
                }
            )
        elif action == "export":
            OfficeExportRequest.model_validate(
                {
                    "output": payload["output"],
                    "version": payload.get("expected_version"),
                }
            )
        elif action == "close":
            assert isinstance(payload["session_id"], str)
    assert set(SPECIALIST_SKILLS) <= seen_skills


class _VisualClient:
    def __init__(self) -> None:
        self.owner: str | None = None
        self.request: OfficeInspectRequest | None = None

    async def inspect(self, request: OfficeInspectRequest, *, owner_session_key: str):
        self.owner = owner_session_key
        self.request = request
        return OfficeInspectSuccess(
            ok=True,
            request_id="inspect_visual",
            session_id=request.session_id,
            version=DocumentVersion(editor_epoch="epoch_visual", model_revision=7),
            result=VisualResult(
                mode="visual",
                data_url="data:image/png;base64,aGVsbG8=",
                width=2,
                height=1,
                target="viewport",
            ),
        )


@pytest.mark.asyncio
async def test_office_tool_returns_versioned_visual_text_and_image_with_owner(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fake = _VisualClient()
    monkeypatch.setattr(
        "mona.agent.tools.office.OfficeServiceClient.from_port",
        lambda _port: fake,
    )
    tool = OfficeTool(workspace=tmp_path, restrict_to_workspace=True)
    tool.set_context(RequestContext(channel="web", chat_id="visual", session_key="owner:visual"))

    result = await tool.execute(
        action="inspect",
        session_id="office_visual",
        query={"mode": "visual"},
    )

    assert isinstance(result, list)
    assert [block["type"] for block in result] == ["text", "image_url"]
    text_payload = json.loads(result[0]["text"])
    assert text_payload["version"] == {
        "editorEpoch": "epoch_visual",
        "modelRevision": 7,
    }
    assert "dataUrl" not in text_payload["result"]
    assert "aGVsbG8=" not in result[0]["text"]
    assert result[1]["image_url"]["url"] == "data:image/png;base64,aGVsbG8="
    assert fake.owner == "owner:visual"
    assert fake.request and fake.request.session_id == "office_visual"


@pytest.mark.parametrize(
    "data_url",
    ["https://example.com/image.png", "data:image/jpeg;base64,aGVsbG8="],
)
def test_visual_result_rejects_external_or_non_png_data_url(data_url: str) -> None:
    with pytest.raises(ValidationError):
        VisualResult.model_validate(
            {
                "mode": "visual",
                "dataUrl": data_url,
                "width": 2,
                "height": 1,
                "target": "viewport",
            }
        )


def test_range_dimensions_and_new_sheet_style_fields_round_trip_through_schema() -> None:
    docs_selection = OfficeInspectResult.model_validate(
        {
            "sessionId": "office_docs",
            "version": {"editorEpoch": "epoch_docs", "modelRevision": 2},
            "result": {
                "mode": "selection",
                "documentType": "docs",
                "blockIds": ["block_1"],
                "text": "已选文本",
            },
        }
    )
    sheets_selection = OfficeInspectResult.model_validate(
        {
            "sessionId": "office_sheet",
            "version": {"editorEpoch": "epoch_sheet", "modelRevision": 4},
            "result": {
                "mode": "selection",
                "documentType": "sheets",
                "sheet": "销售",
                "range": "A1:B2",
            },
        }
    )
    assert docs_selection.result.block_ids == ["block_1"]
    assert docs_selection.result.text == "已选文本"
    assert sheets_selection.result.sheet == "销售"
    assert sheets_selection.result.range == "A1:B2"

    range_result = OfficeInspectResult.model_validate(
        {
            "sessionId": "office_sheet",
            "version": {"editorEpoch": "epoch_sheet", "modelRevision": 4},
            "result": {
                "mode": "range",
                "sheet": "销售",
                "range": "A1:B2",
                "rows": [[{"value": "月份"}, {"value": "收入"}]],
                "columnWidths": [14.0, 18.0],
                "rowHeights": [24.0, 18.0],
            },
        }
    )
    wire_range = range_result.model_dump(by_alias=True)
    assert wire_range["result"]["columnWidths"] == [14.0, 18.0]
    assert wire_range["result"]["rowHeights"] == [24.0, 18.0]

    command = OfficeApplyCommand.model_validate(
        {
            "sessionId": "office_sheet",
            "operationId": "style-1",
            "expectedVersion": {"editorEpoch": "epoch_sheet", "modelRevision": 4},
            "operations": [
                {
                    "op": "set_style",
                    "payload": {
                        "sheet": "销售",
                        "range": "A1:B2",
                        "style": {
                            "fontFamily": "Aptos",
                            "fontSize": 11,
                            "underline": True,
                            "strikethrough": False,
                            "verticalAlign": "center",
                            "wrapText": True,
                            "borderTop": {"style": "thin", "color": "#9DC3E6"},
                            "borderBottom": {"style": "double", "color": "#9DC3E6"},
                            "borderLeft": {"style": "dotted", "color": "#9DC3E6"},
                            "borderRight": {"style": "dashed", "color": "#9DC3E6"},
                        },
                    },
                },
                {
                    "op": "set_column_width",
                    "payload": {"sheet": "销售", "index": 1, "count": 2, "size": 14},
                },
                {
                    "op": "set_row_height",
                    "payload": {"sheet": "销售", "index": 1, "count": 2, "size": 24},
                },
                {
                    "op": "set_page_style",
                    "payload": {"sectionIndex": 0, "widthMm": 210},
                },
                {
                    "op": "set_table_style",
                    "payload": {
                        "blockId": "table_1",
                        "columnWidths": [180, 120],
                        "headerRows": 1,
                        "headerFill": "#D9EAF7",
                        "bodyFill": "#FFFFFF",
                        "borderColor": "#9DC3E6",
                        "cellPadding": 8,
                    },
                },
                {
                    "op": "set_header_footer",
                    "payload": {
                        "kind": "footer",
                        "view": "default",
                        "text": "月度经营报告",
                        "pageNumber": True,
                    },
                },
                {
                    "op": "set_auto_filter",
                    "payload": {"sheet": "销售", "range": "A1:D12"},
                },
                {
                    "op": "set_freeze_panes",
                    "payload": {"sheet": "销售", "rows": 1, "columns": 0},
                },
                {
                    "op": "set_conditional_format",
                    "payload": {
                        "sheet": "销售",
                        "range": "D2:D12",
                        "rule": {"kind": "number", "operator": "lessThan", "value": 0},
                    },
                },
            ],
        }
    )
    wire_command = command.model_dump(by_alias=True)
    style = wire_command["operations"][0]["payload"]["style"]
    assert style["fontFamily"] == "Aptos"
    assert style["fontSize"] == 11.0
    assert style["verticalAlign"] == "center"
    assert style["wrapText"] is True
    assert style["borderBottom"]["style"] == "double"
    assert wire_command["operations"][1]["payload"] == {
        "sheet": "销售",
        "index": 1,
        "count": 2,
        "size": 14,
    }
    assert wire_command["operations"][2]["payload"]["size"] == 24
    assert wire_command["operations"][3]["op"] == "set_page_style"
    assert wire_command["operations"][4]["op"] == "set_table_style"
    assert wire_command["operations"][5]["payload"]["pageNumber"] is True
    assert wire_command["operations"][6]["op"] == "set_auto_filter"
    assert wire_command["operations"][7]["payload"]["rows"] == 1
    assert wire_command["operations"][8]["op"] == "set_conditional_format"


class _CommandSocket:
    def __init__(self) -> None:
        self.closed = False
        self.messages: list[dict[str, object]] = []

    async def send_json(self, message: dict[str, object]) -> None:
        self.messages.append(message)


@pytest.mark.asyncio
async def test_apply_style_defaults_stay_omitted_across_http_and_ws(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    owner = "owner:style"
    command = OfficeApplyCommand.model_validate(
        {
            "sessionId": "office_style",
            "operationId": "style-1",
            "expectedVersion": {"editorEpoch": "epoch_style", "modelRevision": 3},
            "operations": [
                {
                    "op": "set_style",
                    "payload": {
                        "sheet": "销售",
                        "range": "A1",
                        "style": {"bold": True, "borderBottom": None},
                    },
                }
            ],
        }
    )
    captured: dict[str, object] = {}

    async def handler(request: httpx.Request) -> httpx.Response:
        captured["owner"] = request.headers["X-Mona-Session-Key"]
        captured["body"] = json.loads(request.content)
        return httpx.Response(
            200,
            json={
                "ok": True,
                "sessionId": "office_style",
                "operationId": "style-1",
                "version": {"editorEpoch": "epoch_style", "modelRevision": 4},
                "changedTargets": ["销售!A1"],
                "summary": "已更新样式",
            },
        )

    client = OfficeServiceClient(
        "http://127.0.0.1:17174",
        token="token",
        transport=httpx.MockTransport(handler),
    )
    result = await client.apply(command, owner_session_key=owner)

    assert result.ok is True
    assert captured["owner"] == owner
    http_body = captured["body"]
    http_style = http_body["operations"][0]["payload"]["style"]
    assert http_style == {"bold": True, "borderBottom": None}
    forwarded_command = OfficeApplyCommand.model_validate(http_body)

    socket = _CommandSocket()
    hub = OfficeSocketHub()
    await hub.attach("office_style", socket)
    await hub.send_command(forwarded_command)

    ws_style = socket.messages[0]["command"]["operations"][0]["payload"]["style"]
    assert ws_style == {"bold": True, "borderBottom": None}


class _GuardOfficeClient:
    def __init__(self) -> None:
        self.apply_payloads: list[dict[str, object]] = []
        self.inspect_requests: list[OfficeInspectRequest] = []

    async def inspect(self, request: OfficeInspectRequest, *, owner_session_key: str):
        self.inspect_requests.append(request)
        page_index = getattr(request.query, "page_index", None)
        return OfficeInspectSuccess(
            ok=True,
            request_id=f"inspect-{len(self.inspect_requests)}",
            session_id=request.session_id,
            version=DocumentVersion(editor_epoch="epoch_guard", model_revision=1),
            result=VisualResult(
                mode="visual",
                data_url="data:image/png;base64,aGVsbG8=",
                width=2,
                height=1,
                target=str(page_index if page_index is not None else "viewport"),
            ),
        )

    async def apply(self, command: OfficeApplyCommand, *, owner_session_key: str):
        payload = command.model_dump(by_alias=True, mode="json")
        self.apply_payloads.append(payload)
        if "slide-b" in json.dumps(payload):
            return OfficeCommandSuccess(
                ok=True,
                session_id=command.session_id,
                operation_id=command.operation_id,
                version=DocumentVersion(editor_epoch="epoch_guard", model_revision=2),
                changed_targets=["slide-b/element-b"],
                summary="updated new target",
            )
        return OfficeCommandFailure(
            ok=False,
            session_id=command.session_id,
            operation_id=command.operation_id,
            current_version=DocumentVersion(editor_epoch="epoch_guard", model_revision=1),
            changed_targets=[],
            error={
                "code": "INVALID_OPERATION",
                "message": "unsupported property name",
                "retryable": False,
            },
        )


def _apply_call(call_id: str, operation: dict[str, object], *, target: str = "slide-a") -> ToolCallRequest:
    if target == "slide-b":
        operation = {
            "op": "slide_set_text",
            "payload": {"slideId": "slide-b", "elementId": "element-b", "text": "new target"},
        }
    return ToolCallRequest(
        id=call_id,
        name="office",
        arguments={
            "action": "apply",
            "session_id": "office-guard",
            "expected_version": {"editorEpoch": "epoch_guard", "modelRevision": 1},
            "operations": [operation],
        },
    )


@pytest.mark.asyncio
async def test_office_wrong_fields_pause_one_target_but_keep_inspect_and_new_target_available(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fake = _GuardOfficeClient()
    monkeypatch.setattr(
        "mona.agent.tools.office.OfficeServiceClient.from_port",
        lambda _port: fake,
    )
    office = OfficeTool()
    office.set_context(RequestContext(channel="web", chat_id="guard", session_key="owner:guard"))
    tools = ToolRegistry()
    tools.register(office)

    operations = [
        {
            "op": "slide_set_text",
            "payload": {
                "slideId": "slide-a",
                "elementId": "element-a",
                "text": "draft",
                "wrongPropertyOne": True,
            },
        },
        {
            "op": "slide_apply_txn",
            "payload": {
                "ops": [
                    {
                        "op": "set_text",
                        "target": {"slide": "slide-a", "el": "element-a"},
                        "wrongPropertyTwo": True,
                    }
                ]
            },
        },
        {
            "op": "slide_set_font",
            "payload": {
                "slideId": "slide-a",
                "elementId": "element-a",
                "wrongPropertyThree": "red",
            },
        },
        {
            "op": "slide_apply_txn",
            "payload": {
                "ops": [
                    {
                        "op": "set_text",
                        "target": {"slide": "slide-a", "el": "element-a"},
                        "wrongPropertyFour": "blue",
                    }
                ]
            },
        },
        {
            "op": "slide_set_fill",
            "payload": {
                "slideId": "slide-a",
                "elementId": "element-a",
                "wrongPropertyFive": "green",
            },
        },
    ]
    responses: list[LLMResponse] = []
    responses.extend(
        LLMResponse(content="apply", tool_calls=[_apply_call(f"apply-{idx}", operation)])
        for idx, operation in enumerate(operations, start=1)
    )
    for idx in range(1, 5):
        responses.append(
            LLMResponse(
                content="inspect",
                tool_calls=[
                    ToolCallRequest(
                        id=f"inspect-{idx}",
                        name="office",
                        arguments={
                            "action": "inspect",
                            "session_id": "office-guard",
                            "query": {"mode": "visual", "pageIndex": idx},
                        },
                    )
                ],
            )
        )
    # Place the reads around the failing edits and after the target is paused.
    responses = [
        responses[0],
        responses[1],
        responses[5],
        responses[2],
        responses[3],
        responses[6],
        responses[4],
        responses[7],
        LLMResponse(content="apply", tool_calls=[_apply_call("apply-6", operations[0])]),
        responses[8],
        LLMResponse(content="apply", tool_calls=[_apply_call("apply-new", {}, target="slide-b")]),
        LLMResponse(content="done"),
    ]

    provider = MagicMock(spec=LLMProvider)
    provider.chat_with_retry = AsyncMock(side_effect=responses)
    result = await AgentRunner(provider).run(
        AgentRunSpec(
            initial_messages=[],
            tools=tools,
            model="test-model",
            max_iterations=20,
            max_tool_result_chars=50_000,
            repeat_guard_enabled=True,
        )
    )

    assert result.stop_reason == "completed"
    assert result.final_content == "done"
    assert len(fake.apply_payloads) == 6
    assert sum("slide-a" in json.dumps(payload) for payload in fake.apply_payloads) == 5
    assert sum("slide-b" in json.dumps(payload) for payload in fake.apply_payloads) == 1
    assert len(fake.inspect_requests) == 4
    assert [request.query.mode for request in fake.inspect_requests] == ["visual"] * 4
    assert any(event["detail"] == "Office edit target blocked" for event in result.tool_events)
    assert any(event["name"] == "office" and event["status"] == "ok" for event in result.tool_events)
    assert any(
        message.get("role") == "tool"
        and "updated new target" in str(message.get("content"))
        for message in result.messages
    )
    inspect_messages = [
        message
        for message in result.messages
        if message.get("role") == "tool" and isinstance(message.get("content"), list)
    ]
    assert len(inspect_messages) == 4
    assert all(message["content"][1]["type"] == "image_url" for message in inspect_messages)
