from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import pymupdf
import pytest
from docx import Document
from docx.shared import Inches as DocxInches
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as PptxInches

SCRIPT_PATH = (
    Path(__file__).resolve().parents[2]
    / "mona"
    / "skills"
    / "mona-pptx"
    / "scripts"
    / "extract_source_images.py"
)
SPEC = importlib.util.spec_from_file_location("extract_source_images", SCRIPT_PATH)
assert SPEC is not None and SPEC.loader is not None
EXTRACTOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(EXTRACTOR)


@pytest.fixture
def workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("MONA_ACTIVE_WORKSPACE", raising=False)
    return tmp_path


def png_bytes(color: tuple[int, int, int] = (34, 112, 178)) -> bytes:
    image = Image.new("RGB", (80, 40), color)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def create_docx_with_picture(path: Path, picture: bytes) -> None:
    document = Document()
    document.add_paragraph("Quarterly revenue reached the target.")
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(BytesIO(picture), width=DocxInches(1))
    document.add_paragraph("Source: verified sales ledger.")
    document.save(path)


def create_bitmap_pdf(path: Path) -> tuple[float, float, float, float]:
    document = pymupdf.open()
    page = document.new_page(width=600, height=800)
    page.insert_text((80, 80), "Original product image", fontsize=12)
    rect = pymupdf.Rect(80, 100, 280, 200)
    page.insert_image(rect, stream=png_bytes())
    page.insert_text((80, 230), "Image caption text", fontsize=12)
    document.save(path)
    document.close()
    return tuple(rect)


def test_docx_picture_keeps_original_bytes_and_nearby_text(workspace: Path) -> None:
    original = png_bytes()
    create_docx_with_picture(workspace / "source.docx", original)

    result = EXTRACTOR.extract_source_images("source.docx", "assets")

    assert result["ok"] is True
    assert result["source"] == "source.docx"
    assert result["manifestPath"] == "assets/manifest.json"
    assert len(result["assets"]) == 1
    asset = result["assets"][0]
    image_path = workspace / asset["assetPath"]
    assert image_path.is_file()
    assert image_path.read_bytes() == original
    assert asset["sha256"] == hashlib.sha256(original).hexdigest()
    assert (asset["width"], asset["height"]) == (80, 40)
    assert asset["method"] == "embedded"
    occurrence = asset["occurrences"][0]
    assert occurrence["part"] == "word/document.xml"
    assert occurrence["paragraph"] == 2
    assert "Quarterly revenue reached the target." in occurrence["context"]
    assert "Source: verified sales ledger." in occurrence["context"]


def test_docx_external_svg_is_reported_but_never_exported(workspace: Path) -> None:
    source = workspace / "source.docx"
    create_docx_with_picture(source, png_bytes())
    with ZipFile(source) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}

    media_part = next(name for name in members if name.startswith("word/media/"))
    svg_part = media_part.rsplit(".", 1)[0] + ".svg"
    relation_path = "word/_rels/document.xml.rels"
    relation_root = ElementTree.fromstring(members[relation_path])
    image_relation = next(
        relation
        for relation in relation_root
        if relation.get("Type", "").endswith("/image")
        and relation.get("Target") == media_part.removeprefix("word/")
    )
    image_relation.set("Target", svg_part.removeprefix("word/"))
    members[relation_path] = ElementTree.tostring(relation_root, encoding="utf-8", xml_declaration=True)

    content_types_path = "[Content_Types].xml"
    content_types_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    content_types = ElementTree.fromstring(members[content_types_path])
    if not any(node.get("Extension") == "svg" for node in content_types):
        ElementTree.SubElement(
            content_types,
            f"{{{content_types_ns}}}Default",
            {"Extension": "svg", "ContentType": "image/svg+xml"},
        )
    members[content_types_path] = ElementTree.tostring(content_types, encoding="utf-8", xml_declaration=True)

    members.pop(media_part)
    members[svg_part] = (
        b'<svg xmlns="http://www.w3.org/2000/svg" '
        b'xmlns:xlink="http://www.w3.org/1999/xlink" width="80" height="40">'
        b'<image width="80" height="40" xlink:href="https://example.invalid/remote.png"/>'
        b"</svg>"
    )
    svg_source = workspace / "external-svg.docx"
    with ZipFile(svg_source, "w", ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)

    result = EXTRACTOR.extract_source_images("external-svg.docx", "assets")

    assert result["assets"] == []
    assert result["assetCount"] == 0
    assert any("SVG" in warning and "未导出" in warning for warning in result["warnings"])
    manifest = json.loads((workspace / "assets" / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["assets"] == []
    assert any("SVG" in warning and "未导出" in warning for warning in manifest["warnings"])


def test_pptx_records_slides_and_deduplicates_reused_picture(workspace: Path) -> None:
    picture = png_bytes()
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[5])
    first.shapes.title.text = "Customer growth"
    first.shapes.add_picture(BytesIO(picture), PptxInches(1), PptxInches(1))
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = second.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(5), PptxInches(1))
    textbox.text = "Product adoption"
    second.shapes.add_picture(BytesIO(picture), PptxInches(1), PptxInches(2))
    second.shapes.add_picture(BytesIO(picture), PptxInches(4), PptxInches(2))
    presentation.save(workspace / "source.pptx")

    result = EXTRACTOR.extract_source_images("source.pptx", "assets")

    assert len(result["assets"]) == 1
    asset = result["assets"][0]
    image_path = workspace / asset["assetPath"]
    assert image_path.read_bytes() == picture
    occurrences = asset["occurrences"]
    assert [item["slide"] for item in occurrences] == [1, 2, 2]
    assert all(item["part"] == f"ppt/slides/slide{item['slide']}.xml" for item in occurrences)
    assert "Customer growth" in occurrences[0]["context"]
    assert all("Product adoption" in item["context"] for item in occurrences[1:])


def test_pptx_slide_numbers_follow_reordered_presentation_order(workspace: Path) -> None:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[6])
    first_text = first.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(5), PptxInches(1))
    first_text.text = "Original first slide"
    first.shapes.add_picture(BytesIO(png_bytes((220, 40, 40))), PptxInches(1), PptxInches(2))
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second_text = second.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(5), PptxInches(1))
    second_text.text = "Original second slide"
    second.shapes.add_picture(BytesIO(png_bytes((40, 40, 220))), PptxInches(1), PptxInches(2))
    slide_ids = presentation.slides._sldIdLst
    second_id = slide_ids[1]
    slide_ids.remove(second_id)
    slide_ids.insert(0, second_id)
    presentation.save(workspace / "reordered.pptx")

    result = EXTRACTOR.extract_source_images("reordered.pptx", "assets")

    by_context = {
        asset["occurrences"][0]["context"]: asset["occurrences"][0]
        for asset in result["assets"]
    }
    second_occurrence = next(value for key, value in by_context.items() if "Original second slide" in key)
    first_occurrence = next(value for key, value in by_context.items() if "Original first slide" in key)
    assert second_occurrence["slide"] == 1
    assert second_occurrence["part"] == "ppt/slides/slide2.xml"
    assert first_occurrence["slide"] == 2
    assert first_occurrence["part"] == "ppt/slides/slide1.xml"


def test_cli_preview_truncates_repeated_occurrences_but_manifest_is_complete(workspace: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    picture = png_bytes()
    for index in range(6):
        slide.shapes.add_picture(
            BytesIO(picture),
            PptxInches(1 + index % 3),
            PptxInches(1 + index // 3),
        )
    presentation.save(workspace / "reused.pptx")

    completed = subprocess.run(
        [sys.executable, str(SCRIPT_PATH), "--input", "reused.pptx", "--output-dir", "assets"],
        cwd=workspace,
        check=True,
        capture_output=True,
        text=True,
        timeout=30,
    )

    preview = json.loads(completed.stdout)
    assert preview["assetCount"] == 1
    assert preview["truncated"] is True
    assert len(preview["assets"]) == 1
    assert len(preview["assets"][0]["occurrences"]) == 4
    assert preview["assets"][0]["occurrenceCount"] == 6
    manifest = json.loads((workspace / "assets" / "manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["assets"]) == 1
    assert len(manifest["assets"][0]["occurrences"]) == 6


def test_pdf_embedded_bitmap_has_page_and_placement(workspace: Path) -> None:
    expected_rect = create_bitmap_pdf(workspace / "source.pdf")

    result = EXTRACTOR.extract_source_images("source.pdf", "assets", pages=[1])

    assert len(result["assets"]) == 1
    asset = result["assets"][0]
    assert asset["method"] == "embedded"
    image_path = workspace / asset["assetPath"]
    with Image.open(image_path) as image:
        image.verify()
    occurrence = asset["occurrences"][0]
    assert occurrence["page"] == 1
    assert occurrence["bbox"] == pytest.approx(expected_rect)
    assert "Original product image" in occurrence["context"]
    assert "Image caption text" in occurrence["context"]


def test_vector_pdf_region_is_exported_as_nonblank_png(workspace: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=600, height=800)
    page.draw_rect(pymupdf.Rect(100, 110, 300, 210), color=(0.8, 0.1, 0.1), fill=(0.95, 0.75, 0.2))
    page.insert_text((115, 160), "Vector chart", fontsize=22, color=(0.1, 0.2, 0.6))
    page.draw_line((100, 240), (300, 260), color=(0.1, 0.5, 0.2), width=4)
    document.save(workspace / "vector.pdf")
    document.close()

    rect = [70, 80, 330, 280]
    result = EXTRACTOR.extract_source_images("vector.pdf", "crop", pages=[1], rect=rect)

    assert len(result["assets"]) == 1
    asset = result["assets"][0]
    assert asset["method"] == "page_crop"
    assert asset["occurrences"][0]["page"] == 1
    assert asset["occurrences"][0]["bbox"] == pytest.approx(rect)
    image_path = workspace / asset["assetPath"]
    with Image.open(image_path) as image:
        assert image.format == "PNG"
        assert image.width > 0 and image.height > 0
        assert any(low != high for low, high in image.convert("RGB").getextrema())


def test_rotated_pdf_crop_uses_display_coordinates(workspace: Path) -> None:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    native_rect = pymupdf.Rect(40, 40, 180, 180)
    page.draw_rect(native_rect, color=None, fill=(1, 0, 0))
    page.insert_text((65, 75), "TARGET CAPTION", fontsize=8)
    page.set_rotation(90)
    visible_rect = native_rect * page.rotation_matrix
    document.save(workspace / "rotated.pdf")
    document.close()

    crop = pymupdf.Rect(
        visible_rect.x0 + 10,
        visible_rect.y0 + 10,
        visible_rect.x1 - 10,
        visible_rect.y1 - 10,
    )
    result = EXTRACTOR.extract_source_images("rotated.pdf", "rotated-crop", pages=[1], rect=list(crop))

    asset = result["assets"][0]
    assert asset["method"] == "page_crop"
    assert asset["occurrences"][0]["bbox"] == pytest.approx(tuple(crop))
    assert "TARGET CAPTION" in asset["occurrences"][0]["context"]
    with Image.open(workspace / asset["assetPath"]) as image:
        red, green, blue = image.convert("RGB").getpixel((image.width // 2, image.height // 2))
        assert red > 200 and green < 50 and blue < 50


def test_paths_must_stay_in_workspace_and_existing_assets_are_preserved(workspace: Path) -> None:
    create_docx_with_picture(workspace / "source.docx", png_bytes())

    with pytest.raises(ValueError, match="工作区"):
        EXTRACTOR.extract_source_images(str(workspace.parent / "outside.docx"), "assets")
    with pytest.raises(ValueError, match="工作区"):
        EXTRACTOR.extract_source_images("source.docx", str(workspace.parent / "outside-assets"))

    existing = workspace / "existing"
    existing.mkdir()
    marker = existing / "keep.txt"
    marker.write_text("keep this file", encoding="utf-8")
    with pytest.raises(FileExistsError, match="已存在"):
        EXTRACTOR.extract_source_images("source.docx", "existing")
    assert marker.read_text(encoding="utf-8") == "keep this file"
    assert sorted(path.name for path in existing.iterdir()) == ["keep.txt"]


@pytest.mark.parametrize(
    ("pages", "rect", "output_dir"),
    [([2], None, "invalid-page"), ([1], [0, 0, 650, 100], "invalid-rect")],
)
def test_invalid_pdf_page_or_crop_leaves_no_output(
    workspace: Path,
    pages: list[int],
    rect: list[float] | None,
    output_dir: str,
) -> None:
    document = pymupdf.open()
    document.new_page(width=600, height=800)
    document.save(workspace / "source.pdf")
    document.close()

    with pytest.raises(ValueError):
        EXTRACTOR.extract_source_images("source.pdf", output_dir, pages=pages, rect=rect)

    assert not (workspace / output_dir).exists()


def test_cli_uses_active_workspace_when_run_from_skill_directory(workspace: Path) -> None:
    source_name = "原始图片资料.docx"
    output_dir = "extract_source_images.py/引用图片素材"
    create_docx_with_picture(workspace / source_name, png_bytes())
    environment = os.environ.copy()
    environment["MONA_ACTIVE_WORKSPACE"] = str(workspace)

    completed = subprocess.run(
        [sys.executable, str(SCRIPT_PATH), "--input", source_name, "--output-dir", output_dir],
        cwd=SCRIPT_PATH.parent,
        env=environment,
        check=True,
        capture_output=True,
        text=True,
        timeout=30,
    )

    result = json.loads(completed.stdout)
    assert result["ok"] is True
    assert len(result["assets"]) == 1
    assert result["source"] == source_name
    assert result["manifestPath"] == f"{output_dir}/manifest.json"
    assert (workspace / result["manifestPath"]).is_file()
    assert not (SCRIPT_PATH.parent / output_dir).exists()
    assert completed.stderr == ""
