from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = ROOT / "mona" / "skills" / "mona-ppt" / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

from svg_to_pptx.pptx_builder import create_pptx_with_native_svg  # noqa: E402


def _slide_text(slide) -> str:
    chunks: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            chunks.append(shape.text)
    return "\n".join(chunks)


def test_template_underlay_repeats_first_slide_for_svg_export(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    marker = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(4), Inches(0.5))
    marker.text_frame.text = "TEMPLATE_MARKER"
    prs.save(template)

    svg_files: list[Path] = []
    for index in range(1, 3):
        svg = tmp_path / f"slide_{index}.svg"
        svg.write_text(
            f"""<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">
  <rect x="180" y="160" width="600" height="220" fill="#ffffff"/>
  <text x="220" y="260" font-size="44" fill="#111111">CONTENT_{index}</text>
</svg>""",
            encoding="utf-8",
        )
        svg_files.append(svg)

    output = tmp_path / "out.pptx"
    ok = create_pptx_with_native_svg(
        svg_files,
        output,
        verbose=False,
        use_native_shapes=True,
        template_underlay_pptx=template,
    )

    assert ok
    out = Presentation(str(output))
    assert len(out.slides) == 2
    assert "TEMPLATE_MARKER" in _slide_text(out.slides[0])
    assert "TEMPLATE_MARKER" in _slide_text(out.slides[1])
    assert "CONTENT_1" in _slide_text(out.slides[0])
    assert "CONTENT_2" in _slide_text(out.slides[1])
